# -*- coding: utf-8 -*-
"""
接口自动化回归测试执行器
====================================
基于主流程测试用例(Excel)映射的接口调用链, 自动执行接口回归测试。

两种运行模式:
  1) verify 模式(默认, 无需Token):
     对每个接口发送未认证探测请求, 验证路由存在性与认证网关正常
     (401=>存在需认证, 200=>可达, 503=>被禁用, 404=>接口已下线)
  2) full 模式(需要 Token):
     携带真实 Token 执行完整业务流程, 断言业务响应 (2xx/4xx语义)

用法:
  python -m api_test.regression_runner --mode verify
  python -m api_test.regression_runner --mode full --token sk-xxx
  python -m api_test.regression_runner --mode full --token-file token.txt
  python -m api_test.regression_runner --case Agent-012 --mode verify   # 单条

输出:
  output/回归报告/接口回归报告_<时间戳>.xlsx  (Excel 明细)
  output/回归报告/接口回归报告_<时间戳>.md    (Markdown 摘要)
"""
import argparse
import json
import os
import re
import ssl
import sys
import time
import urllib.parse
import urllib.request
import urllib.error
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.console import console
from api_test.regression_scenarios import load_all, SERVICE_BASE

OUT_DIR = ROOT / "output" / "回归报告"


# ============================================================
# HTTP 工具 (基于 urllib, 零额外依赖)
# ============================================================
_CTX = ssl.create_default_context()
_CTX.check_hostname = False
_CTX.verify_mode = ssl.CERT_NONE


def send_request(method: str, url: str, headers: dict, body=None, timeout: int = 15):
    """发送 HTTP 请求, 返回 (status, headers, body_text)"""
    data = None
    if body is not None:
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        headers.setdefault("Content-Type", "application/json")
    req = urllib.request.Request(url, data=data, headers=headers, method=method.upper())
    try:
        resp = urllib.request.urlopen(req, timeout=timeout, context=_CTX)
        text = resp.read().decode("utf-8", "ignore")
        return resp.status, dict(resp.headers), text
    except urllib.error.HTTPError as e:
        text = e.read().decode("utf-8", "ignore")
        return e.code, dict(e.headers), text
    except Exception as e:
        return -1, {}, f"ERR: {str(e)[:200]}"


def fill_path(path: str, ctx: dict, case_id: str, service: str = "dify") -> str:
    """填充路径变量 {id}/{task_id} 等 (service=ragflow 时使用 ragflow 侧 ID)"""
    out = path
    ds_id = ctx.get("rf_dataset_id") if service == "ragflow" else ctx.get("dataset_id", "")
    ds_default = ds_id or "01234567-89ab-4def-8123-456789abcdef"
    out = out.replace("{id}", ctx.get("app_id") or "01234567-89ab-4def-8123-456789abcdef")
    out = out.replace("{task_id}", ctx.get("task_id", "01234567-89ab-4def-8123-456789abcdef"))
    out = out.replace("{dataset_id}", ds_default)
    out = out.replace("{document_id}", ctx.get("document_id", "01234567-89ab-4def-8123-456789abcdef"))
    out = out.replace("{doc_id}", ctx.get("document_id", "01234567-89ab-4def-8123-456789abcdef"))
    out = out.replace("{chunk_id}", ctx.get("chunk_id", "01234567-89ab-4def-8123-456789abcdef"))
    out = out.replace("{tag_id}", ctx.get("tag_id", "01234567-89ab-4def-8123-456789abcdef"))
    out = out.replace("{tag_base_id}", ctx.get("tag_base_id") or "01234567-89ab-4def-8123-456789abcdef")
    out = out.replace("{provider}", ctx.get("provider", "openai"))
    out = out.replace("{model_type}", ctx.get("model_type", "text-generation"))
    out = re.sub(r"\{[a-zA-Z_]+\}", "01234567-89ab-4def-8123-456789abcdef", out)
    return out


# ============================================================
# 回归执行器
# ============================================================
class RegressionRunner:
    def __init__(self, token: str = None, ragflow_token: str = None,
                 mode: str = "verify", timeout: int = 15,
                 delay: float = 0.1, readonly: bool = False):
        self.token = token
        self.ragflow_token = ragflow_token or token    # RagFlow 域可单独提供
        self.mode = mode          # verify | full
        self.timeout = timeout
        self.delay = delay
        self.readonly = readonly  # full模式中仅执行只读步骤(GET/检索), 跳过写操作
        self.results = []         # 每条用例的详细结果
        self.env = {}             # 上下文变量(在full模式中传递ID)

    # 判定某步骤是否只读 (写方法 + 非检索路径 = 写操作)
    def _is_readonly_step(self, step: dict) -> bool:
        method = step.get("method", "GET").upper()
        if method == "GET":
            return True
        if method == "POST":
            path = step.get("path", "")
            # 检索/测试类 POST 视为只读
            if any(k in path for k in ("hit-testing", "retrieval-test", "retrieval_test",
                                       "use-check", "check-dependencies", "dynamic-options",
                                       "latest-versions")):
                return True
            return False
        return False

    # ---------- 认证头 ----------
    def _headers(self, service: str, step: dict) -> dict:
        h = {
            "User-Agent": "Mozilla/5.0 (Regression Runner)",
            "Accept": "application/json",
        }
        tok = self.ragflow_token if service == "ragflow" else self.token
        if tok:
            h["Authorization"] = f"Bearer {tok}"
        return h

    # ---------- full 模式: 预取真实资源 ID (避免假UUID404) ----------
    def prefetch_real_ids(self):
        """从真实环境读取 app_id / dataset_id 填入 env, 供 {id} 占位替换"""
        if self.mode != "full" or not (self.token or self.ragflow_token):
            return
        base = SERVICE_BASE["dify"]
        # 应用列表 -> 优先选 advanced-chat / workflow 模式(有 workflow 草稿)
        try:
            code, _, body = send_request("GET", base + "/console/api/apps", {
                "Authorization": f"Bearer {self.token}",
                "Accept": "application/json",
            }, None, 12)
            if code == 200:
                data = json.loads(body)
                items = data.get("data", [])
                if items:
                    pref = next((a for a in items if a.get("mode") in ("advanced-chat", "workflow")), items[0])
                    self.env["app_id"] = pref.get("id", "")
                    self.env["app_mode"] = pref.get("mode", "")
        except Exception:
            pass
        # 数据集列表 -> 第一个 dataset
        try:
            code, _, body = send_request("GET", base + "/console/api/datasets", {
                "Authorization": f"Bearer {self.token}",
                "Accept": "application/json",
            }, None, 12)
            if code == 200:
                data = json.loads(body)
                items = data.get("data", [])
                if items:
                    self.env["dataset_id"] = items[0].get("id", "")
        except Exception:
            pass
        # RagFlow 数据集 + 文档
        try:
            code, _, body = send_request("GET", SERVICE_BASE["ragflow"] + "/api/v1/datasets", {
                "Authorization": f"Bearer {self.ragflow_token}",
                "Accept": "application/json",
            }, None, 12)
            if code == 200:
                data = json.loads(body)
                items = data.get("data", [])
                if items:
                    self.env["rf_dataset_id"] = items[0].get("id", "")
                    # 预取第一个文档ID (供 documents/{document_id}/... 步骤)
                    try:
                        code2, _, body2 = send_request(
                            "GET", SERVICE_BASE["ragflow"] + f"/api/v1/datasets/{items[0].get('id')}/documents?page=1&page_size=1",
                            {"Authorization": f"Bearer {self.ragflow_token}", "Accept": "application/json"},
                            None, 12)
                        if code2 == 200:
                            d2 = json.loads(body2)
                            docs = ((d2.get("data") or {}).get("docs") or []) or []
                            if docs:
                                self.env["document_id"] = docs[0].get("id", "")
                    except Exception:
                        pass
        except Exception:
            pass

    # ---------- 单步执行 ----------
    def _run_step(self, case_id: str, case: dict, step: dict) -> dict:
        service = step.get("service") or case.get("service", "dify")
        base = step.get("base", "") or SERVICE_BASE.get(service, "http://ai-func.ibosssoft.com.cn")
        # full 模式中替换 ${EMAIL}/${PASSWORD} 环境变量
        path = fill_path(step["path"], self.env, case_id, service)
        url = base.rstrip("/") + (path if path.startswith("/") else "/" + path)

        headers = self._headers(service, step)
        body = None
        if self.mode == "full":
            body = step.get("params") or {}
            # 替换 ${...} 占位
            body = self._resolve_placeholders(body, case_id)

        t0 = time.time()
        status, resp_headers, resp_text = send_request(
            step.get("method", "GET"), url, headers, body, self.timeout)
        cost = round(time.time() - t0, 2)

        # 判定
        # verify 模式语义: 验证"接口路由是否存在/网关是否正常"
        #   401/403 = 路由存在需认证  405 = 路由存在方法错  503 = 存在但禁用
        #   200-399 = 可达           404 = 路由不存在      -1=网络异常
        expect = step.get("expect") or ["200"]
        if self.mode == "verify":
            if status == 404 or status == -1:
                passed = False
                verdict = "FAIL"
                detail = f"路由不存在/网络异常 实得{status}"
            elif status in (401, 403) or (200 <= status < 400) or status in (405, 503):
                passed = True
                verdict = "PASS"
                note = {401: "存在需认证", 403: "存在禁止访问", 405: "存在方法待修正",
                        503: "存在功能禁用"}.get(status, "")
                detail = f"实得{status} 路由存在{('(' + note + ')') if note else ''}"
            else:
                passed = False
                verdict = "FAIL"
                detail = f"意外状态 {status}"
        else:
            # full 模式语义:
            #   503 = 环境功能禁用 -> BLOCKED(受限, 非接口缺陷)
            #   404 = 资源/路由不存在 -> FAIL
            ec = step.get("expect_code", 200)
            expect_any = step.get("expect_any", [])
            msg = (resp_text or "").strip()[:150].replace("\n", " ")
            # RagFlow 域: HTTP 200 但业务 code 非 0 时按业务语义判定
            biz_code = None
            if service == "ragflow" and resp_text:
                try:
                    _d = json.loads(resp_text)
                    if isinstance(_d, dict) and isinstance(_d.get("code"), (int, str)):
                        biz_code = str(_d.get("code"))
                except Exception:
                    pass
            if service == "ragflow" and biz_code is not None:
                if biz_code == "0" or biz_code == 0 or biz_code == "200":
                    verdict = "PASS"
                    detail = f"业务成功(code={biz_code})"
                    passed = True
                elif biz_code in ("109", "102"):
                    verdict = "BLOCKED"
                    detail = f"路由存在但凭证/资源受限(code={biz_code}): {msg}"
                    passed = False
                elif biz_code == "101":
                    verdict = "BLOCKED"
                    detail = f"路由存在但参数缺失(code=101): {msg}"
                    passed = False
                elif biz_code == "100":
                    if "404 Not Found" in msg:
                        verdict = "FAIL"
                        detail = f"路由不存在(code=100/404): {msg}"
                    elif "405" in msg:
                        verdict = "FAIL"
                        detail = f"路由方法错误(code=100/405): {msg}"
                    elif status in expect_any or (expect_any and status in expect_any):
                        verdict = "PASS"
                        detail = f"实得{status} {('[' + step.get('note_hint', '') + ']') if step.get('note_hint') else ''} {msg}"
                        passed = True
                    else:
                        verdict = "FAIL"
                        detail = f"业务错误(code=100): {msg}"
                        passed = False
                else:
                    verdict = "FAIL"
                    detail = f"业务错误(code={biz_code}) {msg}"
                    passed = False
            elif status == 503:
                passed = False
                verdict = "BLOCKED"
                detail = f"环境功能禁用(503)"
            elif status == 404:
                passed = False
                verdict = "FAIL"
                detail = f"资源/路由不存在(404) 期望{ec}"
            elif expect_any and status in expect_any:
                passed = True
                verdict = "PASS"
                detail = f"实得{status}(预期之一)"
                if status != 200:
                    hint = step.get("note_hint", "")
                    detail = f"实得{status} {('[' + hint + ']') if hint else ''} {msg}"
            else:
                passed = (status == ec)
                verdict = "PASS" if passed else "FAIL"
                detail = f"期望{ec} 实得{status}"
                if not passed and msg:
                    detail += f" | {msg}"

        return {
            "case_id": case_id,
            "case_title": case.get("title", ""),
            "module": case.get("module", ""),
            "priority": case.get("priority", ""),
            "service": service,
            "base": base,
            "method": step.get("method", "GET"),
            "path": path,
            "expect": expect,
            "actual_status": status,
            "verdict": verdict,
            "detail": detail,
            "cost_s": cost,
            "resp_body": resp_text[:200] if resp_text else "",
        }

    def _resolve_placeholders(self, body, case_id):
        """递归替换 ${VAR} (取环境变量, 缺失保留占位)"""
        if isinstance(body, dict):
            return {k: self._resolve_placeholders(v, case_id) for k, v in body.items()}
        if isinstance(body, list):
            return [self._resolve_placeholders(v, case_id) for v in body]
        if isinstance(body, str):
            def rep(m):
                name = m.group(1)
                return os.environ.get(name, m.group(0))
            return re.sub(r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}", rep, body)
        return body

    # ---------- 整条用例执行 ----------
    def run_case(self, case_id: str, case: dict) -> dict:
        steps = case.get("steps", [])
        if not steps:
            return {"case_id": case_id, "case_title": case.get("title", ""),
                    "module": case.get("module", ""), "priority": case.get("priority", ""),
                    "tag": case.get("tag", "UI"), "verdict": "SKIP",
                    "detail": "纯UI用例,未映射接口", "steps": []}

        step_results = []
        for step in steps:
            # readonly 模式: 跳过写操作步骤
            if self.mode == "full" and self.readonly and not self._is_readonly_step(step):
                step_results.append({
                    "case_id": case_id, "case_title": case.get("title", ""),
                    "module": case.get("module", ""), "priority": case.get("priority", ""),
                    "service": step.get("service") or case.get("service", "dify"),
                    "base": step.get("base", ""), "method": step.get("method", "GET"),
                    "path": step.get("path", ""), "expect": [], "actual_status": "-",
                    "verdict": "SKIP", "detail": "readonly模式跳过写操作",
                    "cost_s": 0, "resp_body": "",
                })
                continue
            r = self._run_step(case_id, case, step)
            step_results.append(r)
            if self.delay:
                time.sleep(self.delay)
            # 上下文: full模式记录 app_id / dataset_id / document_id
            if self.mode == "full" and r["actual_status"] in (200, 201):
                self._capture_ids(r)

        # 判定: 跳过步骤不计失败; 若所有步骤均被跳过则用例判 SKIP
        executed = [r for r in step_results if r["verdict"] != "SKIP"]
        if not executed:
            verdict = "SKIP"
        elif any(r["verdict"] == "FAIL" for r in executed):
            verdict = "FAIL"
        elif all(r["verdict"] == "PASS" for r in executed):
            verdict = "PASS"
        else:
            verdict = "BLOCKED"   # 有步骤受限(503环境禁用)
        fail_n = sum(1 for r in executed if r["verdict"] == "FAIL")
        blocked_n = sum(1 for r in step_results if r["verdict"] == "BLOCKED")
        skip_n = sum(1 for r in step_results if r["verdict"] == "SKIP")
        detail = f"{len(executed)}步, 失败{fail_n}步"
        if blocked_n:
            detail += f", 环境禁用{blocked_n}步"
        if skip_n:
            detail += f", 跳过{skip_n}步"
        return {
            "case_id": case_id, "case_title": case.get("title", ""),
            "module": case.get("module", ""), "priority": case.get("priority", ""),
            "tag": case.get("tag", ""), "note": case.get("note", ""),
            "verdict": verdict,
            "detail": detail,
            "steps": step_results,
        }

    def _capture_ids(self, r: dict):
        """从成功响应中提取 id 供后续步骤使用"""
        body = r.get("resp_body", "")
        try:
            data = json.loads(body)
        except Exception:
            return
        for key in ("id", "app_id", "dataset_id", "document_id", "task_id", "chunk_id", "tag_id"):
            if key in data and isinstance(data[key], str):
                self.env[key] = data[key]

    # ---------- 全量执行 ----------
    def run_all(self, scenarios: dict, only: str = None, prefix: str = None) -> list:
        results = []
        items = [(k, v) for k, v in scenarios.items()]
        if prefix:
            items = [it for it in items if it[0].startswith(prefix)]
        if only:
            items = [it for it in items if it[0] == only]
        for case_id, case in items:
            r = self.run_case(case_id, case)
            results.append(r)
            mark = {"PASS": "✅", "FAIL": "❌", "SKIP": "⏭️", "BLOCKED": "🚫"}.get(r["verdict"], "❓")
            console.print(f"  {mark} [{r['verdict']:4s}] {case_id} {r['case_title']} - {r['detail']}")
        return results

    # ---------- 报告 ----------
    def generate_report(self, results: list, scenarios: dict) -> dict:
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_xlsx = OUT_DIR / f"接口回归报告_{ts}.xlsx"
        out_md = OUT_DIR / f"接口回归报告_{ts}.md"

        # --- 统计 ---
        vc = Counter(r["verdict"] for r in results)
        step_total = sum(len(r["steps"]) for r in results)
        step_fail = sum(1 for r in results for s in r["steps"] if s["verdict"] == "FAIL")
        pc = Counter(r["priority"] for r in results)

        # --- MD ---
        md = []
        md.append(f"# 接口自动化回归测试报告")
        md.append("")
        md.append(f"**执行时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        md.append(f"**运行模式**: {self.mode}" + (" (携带Token: " + self.token[:8] + "...)" if self.token else ""))
        md.append(f"**用例总数**: {len(results)}  通过 {vc.get('PASS',0)} / 失败 {vc.get('FAIL',0)} / 跳过 {vc.get('SKIP',0)}")
        md.append(f"**接口步骤**: {step_total} 步, 失败 {step_fail} 步")
        md.append(f"**优先级**: {dict(pc)}")
        md.append("")
        md.append("## 结果明细")
        md.append("")
        md.append("| 用例 | 模块 | 优先级 | 结果 | 详情 |")
        md.append("|---|---|---|---|---|")
        for r in results:
            badge = {"PASS": "✅", "SKIP": "⏭️", "BLOCKED": "🚫", "FAIL": "❌"}.get(r["verdict"], "❓")
            md.append(f"| {r['case_id']} | {r['module']} | {r['priority']} | {badge} {r['verdict']} | {r['detail']} |")
        md.append("")
        md.append("## 接口步骤明细 (失败项)")
        md.append("")
        for r in results:
            for s in r["steps"]:
                if s["verdict"] == "FAIL":
                    note = f"  ({r.get('note', '')})" if r.get("note") else ""
                    md.append(f"- **{s['case_id']}** {s['method']} `{s['path']}` → 实得 {s['actual_status']}{note}")
        md.append("")
        md.append("## 用例备注/已知缺口")
        md.append("")
        for r in results:
            if r.get("note"):
                md.append(f"- **{r['case_id']}**: {r['note']}")
        md.append("")
        md.append("## 说明")
        md.append("")
        md.append("- verify 模式验证路由/网关可达性; full 模式需 `--token` 携带有效凭证验证业务流程。")
        md.append("- 401 = 路由存在但需认证; 404 = 接口下线; 503 = 功能被禁用。")
        out_md.write_text("\n".join(md), encoding="utf-8")

        # --- XLSX ---
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "回归结果"
        hdrs = ["用例编号", "用例标题", "模块", "优先级", "结果", "详情"]
        for c, h in enumerate(hdrs, 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")

        fills = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "SKIP": "EDEDED", "BLOCKED": "FFEB9C"}
        for r_i, r in enumerate(results, 2):
            for c, val in enumerate([r["case_id"], r["case_title"], r["module"],
                                     r["priority"], r["verdict"], r["detail"]], 1):
                cell = ws.cell(r_i, c, val)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            fcol = fills.get(r["verdict"])
            if fcol:
                ws.cell(r_i, 5).fill = PatternFill("solid", fgColor=fcol)

        ws2 = wb.create_sheet("接口步骤")
        hdrs2 = ["用例", "服务", "方法", "路径", "期望", "实得", "结果", "耗时s", "响应"]
        for c, h in enumerate(hdrs2, 1):
            cell = ws2.cell(1, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        i = 2
        for r in results:
            for s in r["steps"]:
                for c, val in enumerate([s["case_id"], s["service"], s["method"], s["path"],
                                         ",".join(map(str, s["expect"])), s["actual_status"],
                                         s["verdict"], s["cost_s"], s["resp_body"]], 1):
                    cell = ws2.cell(i, c, val)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                if s["verdict"] == "FAIL":
                    ws2.cell(i, 7).fill = PatternFill("solid", fgColor="FFC7CE")
                i += 1

        ws2.auto_filter.ref = f"A1:I{i-1}"
        ws.auto_filter.ref = f"A1:F{len(results)+1}"
        for col, w in zip("ABCDEF", [12, 34, 18, 8, 8, 28]):
            ws.column_dimensions[col].width = w
        for col, w in zip("ABCDEFGHI", [12, 10, 8, 52, 10, 8, 8, 8, 50]):
            ws2.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws2.freeze_panes = "A2"

        wb.save(out_xlsx)
        return {"xlsx": str(out_xlsx), "md": str(out_md),
                "summary": {"total": len(results), "pass": vc.get("PASS", 0),
                            "fail": vc.get("FAIL", 0), "skip": vc.get("SKIP", 0),
                            "steps": step_total, "step_fail": step_fail}}


# ============================================================
# CLI
# ============================================================
def main():
    parser = argparse.ArgumentParser(description="接口自动化回归测试")
    parser.add_argument("--mode", choices=["verify", "full"], default="verify",
                        help="verify=路由可达性(无需Token); full=业务回归(需Token)")
    parser.add_argument("--token", default=None, help="Dify Bearer Token(环境变量 TOKEN 亦可)")
    parser.add_argument("--token-ragflow", default=None,
                        help="RagFlow API Key(如与--token不同; 缺省复用--token)")
    parser.add_argument("--token-file", default=None, help="从文件读取 Token")
    parser.add_argument("--case", default=None, help="仅执行单条用例(如 Agent-012)")
    parser.add_argument("--prefix", default=None, help="按用例编号前缀过滤(如 --prefix RAG)")
    parser.add_argument("--timeout", type=int, default=15, help="请求超时秒")
    parser.add_argument("--delay", type=float, default=0.1, help="请求间隔秒")
    parser.add_argument("--readonly", action="store_true",
                        help="full模式下仅执行只读步骤(GET/检索), 跳过写操作(安全)")
    args = parser.parse_args()

    token = args.token or os.environ.get("TOKEN") or os.environ.get("LLM_API_KEY")
    if args.token_file:
        token = Path(args.token_file).read_text(encoding="utf-8").strip()
    ragflow_token = args.token_ragflow or args.token or os.environ.get("RAGFLOW_API_KEY")

    scenarios, meta = load_all()

    console.print("=" * 70)
    mode_label = "verify 路由验证"
    if args.mode == "full":
        mode_label = "full 业务回归" + ("(只读安全模式)" if args.readonly else "(含写操作)")
    console.print(f"[bold]接口自动化回归测试 - {mode_label}[/bold]")
    console.print(f"  场景总数: {len(scenarios)}  接口清单: {meta['platform_yaml']}")
    if token:
        console.print(f"  Dify Token: {token[:10]}... (已提供)")
    if ragflow_token:
        console.print(f"  RagFlow Key: {ragflow_token[:10]}... (已提供)")
    console.print("=" * 70)

    runner = RegressionRunner(token=token, ragflow_token=ragflow_token,
                              mode=args.mode,
                              timeout=args.timeout, delay=args.delay,
                              readonly=args.readonly)
    # full 模式: 预取真实 app_id/dataset_id, 替换路径占位
    if args.mode == "full" and (token or ragflow_token):
        console.print("[dim]预取真实资源ID (apps/datasets)...[/dim]")
        runner.prefetch_real_ids()
        if runner.env.get("app_id") or runner.env.get("dataset_id") or runner.env.get("rf_dataset_id"):
            console.print(f"[dim]  预取成功: app={runner.env.get('app_id','-')[:8]}... "
                          f"dataset={runner.env.get('dataset_id','-')[:8]}... "
                          f"ragflow_dataset={runner.env.get('rf_dataset_id','-')[:8]}...[/dim]")
        else:
            console.print("[dim]  预取未获得ID, 将使用占位UUID[/dim]")
    results = runner.run_all(scenarios, only=args.case, prefix=args.prefix)

    report = runner.generate_report(results, scenarios)

    console.print("")
    console.print("[bold]执行统计[/bold]")
    s = report["summary"]
    console.print(f"  用例: {s['total']}  通过 {s['pass']}  失败 {s['fail']}  跳过 {s['skip']}")
    console.print(f"  接口步骤: {s['steps']} 失败 {s['step_fail']}")
    console.print(f"[bold green]报告: {report['xlsx']}[/bold green]")
    console.print(f"[bold green]摘要: {report['md']}[/bold green]")
    return 0


if __name__ == "__main__":
    sys.exit(main())