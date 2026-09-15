#!/usr/bin/env python3
"""验证全量 API 用例 Excel"""
import sys, glob, os
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

OUT = ROOT / "output" / "api_full_check.txt"
lines = []
def log(m=""):
    lines.append(m)

from openpyxl import load_workbook

files = sorted(glob.glob(str(ROOT / "output" / "*全量*API用例*.xlsx")), key=os.path.getmtime)
if not files:
    log("未找到全量 API 用例 Excel")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("done")
    sys.exit(0)

f = files[-1]
log(f"FILE: {f}")
log(f"SIZE: {os.path.getsize(f)} bytes")
wb = load_workbook(f)
log(f"SHEETS: {wb.sheetnames}")
ws = wb["API测试用例"]
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
log(f"COLS: {ws.max_column}  ROWS: {ws.max_row - 1}")
log(f"HEADERS: {hdr}")
log("")

required = ["序号", "用例编号", "接口", "接口名称", "HTTP方法", "优先级",
            "用例标题", "前置条件", "请求头", "请求参数", "操作步骤",
            "预期结果", "实际结果", "备注", "关联需求"]
empty_count = {h: 0 for h in required}
for r in range(2, ws.max_row + 1):
    for c, h in enumerate(required, 1):
        v = ws.cell(r, c).value
        if v is None or str(v).strip() == "":
            empty_count[h] += 1

log("=== 空值统计 ===")
for h, n in empty_count.items():
    flag = "OK(执行时填写)" if h == "实际结果" else ("!! 有空值" if n else "OK")
    log(f"  {h}: {n} 空  {flag}")

log("")
log("=== 方法分布 ===")
from collections import Counter
mc = Counter()
for r in range(2, ws.max_row + 1):
    mc[str(ws.cell(r, 5).value)] += 1
for m, n in mc.most_common():
    log(f"  {m}: {n}")

log("")
log("=== 优先级分布 ===")
pc = Counter()
for r in range(2, ws.max_row + 1):
    pc[str(ws.cell(r, 6).value)] += 1
for p, n in sorted(pc.items()):
    log(f"  {p}: {n}")

log("")
log("=== 接口覆盖数 ===")
ifaces = set()
for r in range(2, ws.max_row + 1):
    ifaces.add(str(ws.cell(r, 3).value))
log(f"  覆盖接口数: {len(ifaces)}")

log("")
log("=== 抽样(行2, 150, 300, 473) ===")
for r in [2, 150, 300, 473]:
    if r > ws.max_row:
        continue
    log(f"--- 第 {r} 行 ---")
    for c, h in enumerate(required, 1):
        v = str(ws.cell(r, c).value or "").replace("\n", " ⏎ ")
        log(f"  {h}: {v[:160]}")
    log("")

OUT.write_text("\n".join(lines), encoding="utf-8")
print("done")