#!/usr/bin/env python3
"""完整读取两个主流程 Excel 全部用例"""
import sys, json
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

OUT = ROOT / "output" / "mainflow_full.json"
ALL = []

from openpyxl import load_workbook

files = [
    ("agent", r"C:\Users\Admin\Desktop\需求\主流程\智能体平台开发态-主流程回归测试用例.xlsx"),
    ("rag", r"C:\Users\Admin\Desktop\需求\主流程\知识库平台-主流程回归测试用例.xlsx"),
]

for tag, f in files:
    wb = load_workbook(f, data_only=True)
    ws = wb["功能测试用例"]
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, 10)]
        if not any(v is not None and str(v).strip() for v in row):
            continue
        ALL.append({
            "platform": tag,
            "序号": row[0], "用例编号": row[1], "模块": row[2], "优先级": row[3],
            "类型": row[4], "用例标题": row[5], "前置条件": row[6],
            "操作步骤": row[7], "实际结果": row[8],
        })

OUT.write_text(json.dumps(ALL, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"共 {len(ALL)} 条用例 -> {OUT}")
from collections import Counter
print("模块分布:", dict(Counter(f"{x['platform']}:{x['模块']}" for x in ALL)))
print("优先级分布:", dict(Counter(x["优先级"] for x in ALL)))