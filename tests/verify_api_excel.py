#!/usr/bin/env python3
"""验证 API 用例 Excel 质量"""
import sys, glob, os
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

OUT = ROOT / "output" / "api_case_check.txt"
lines = []
def log(m=""):
    lines.append(m)

from openpyxl import load_workbook

files = sorted(glob.glob(str(ROOT / "output" / "*API用例*.xlsx")), key=os.path.getmtime)
if not files:
    log("未找到 API 用例 Excel")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("done")
    sys.exit(0)

f = files[-1]
log(f"FILE: {f}")
wb = load_workbook(f)
log(f"SHEETS: {wb.sheetnames}")
ws = wb["API测试用例"]
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
log(f"COLS: {ws.max_column}  ROWS: {ws.max_row - 1}")
log(f"HEADERS: {hdr}")
log("")

# 完整性检查
required = ["序号", "用例编号", "接口", "接口名称", "HTTP方法", "优先级",
            "用例标题", "前置条件", "请求头", "请求参数", "操作步骤",
            "预期结果", "实际结果", "备注", "关联需求"]
empty_count = {h: 0 for h in required}
for r in range(2, ws.max_row + 1):
    for c, h in enumerate(required, 1):
        v = ws.cell(r, c).value
        if v is None or str(v).strip() == "":
            empty_count[h] += 1

log("=== 空值统计 ===")
for h, n in empty_count.items():
    flag = "OK(执行时填写)" if h == "实际结果" else ("!! 有空值" if n else "OK")
    log(f"  {h}: {n} 空  {flag}")

# 抽样
log("")
log("=== 抽样用例 ===")
for r in [2, 7, 12, 18]:
    if r > ws.max_row:
        continue
    log(f"--- 第 {r} 行 ---")
    for c, h in enumerate(required, 1):
        v = str(ws.cell(r, c).value or "").replace("\n", " ⏎ ")
        log(f"  {h}: {v[:200]}")
    log("")

# 摘要 sheet
if "摘要" in wb.sheetnames:
    ws2 = wb["摘要"]
    log("=== 摘要 ===")
    for r in range(1, ws2.max_row + 1):
        vals = [str(ws2.cell(r, c).value or "") for c in range(1, 3)]
        if any(v for v in vals):
            log("  " + " | ".join(v for v in vals if v))

OUT.write_text("\n".join(lines), encoding="utf-8")
print("done")