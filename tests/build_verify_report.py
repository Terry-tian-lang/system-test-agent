#!/usr/bin/env python3
"""
生成接口验证综合报告 (Excel + Markdown):
- 列出全部111个接口及验证状态
- 汇总三服务架构的验证结论
"""
import sys, json, yaml
from pathlib import Path
from datetime import datetime

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

OUT_MD = ROOT / "output" / "接口验证综合报告.md"
OUT_XLSX = ROOT / "output" / "接口验证报告.xlsx"
lines = []
def log(m=""):
    lines.append(m)

# 读取合并清单
merged = yaml.safe_load((ROOT / "api_test" / "interfaces" / "ai_func_platform_api.yaml").read_text(encoding="utf-8"))

# 读取原始验证结果
verify = json.loads((ROOT / "output" / "api_verify_result.json").read_text(encoding="utf-8"))
vmap = {v["path"]: v for v in verify}

# ---------- 汇总统计 ----------
svc_rows = []  # (服务, 分组, 路径, 方法, 验证状态, 说明)
for svc in ["dify_console", "dify_runtime", "ragflow", "cas"]:
    items = merged.get(svc, [])
    if not isinstance(items, list):
        continue
    group = svc
    for it in items:
        if not isinstance(it, dict) or "path" not in it:
            continue
        p = it["path"]
        verified = it.get("verified", "UNKNOWN")
        methods = it.get("methods", "try")
        note = {"VALID": "实测:未经认证返回401，路由存在",
                "METHOD": "实测:方法待修正",
                "DISABLED": "实测:功能被禁用(503)",
                "OPEN": "实测:可公开访问",
                "VALID-实测401": "实测:未经认证返回401",
                "VALID-代码探测": "依据开源代码推测，需登录确认"}.get(verified, verified)
        svc_rows.append((svc, group, p, methods, verified, note))

# ---------- Markdown 报告 ----------
md = []
md.append("# 接口验证综合报告 — ibosssoft AI 功能平台")
md.append("")
md.append(f"**验证时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
md.append("")
md.append("## 一、平台架构（三服务确认）")
md.append("")
md.append("| 服务 | 类型 | 地址 | 认证方式 | 验证证据 |")
md.append("|---|---|---|---|---|")
md.append("| ai-func | **魔改版 Dify** (Agent/应用平台) | http://ai-func.ibosssoft.com.cn | Bearer Token | 未认证 401 `{'code':'unauthorized'}`；`/runtime/console/api/*` 独立返回 `Invalid token.` |")
md.append("| rag-func | **魔改版 RagFlow** (知识库/RAG) | http://rag-func.ibosssoft.com.cn | Authorization Header (API Key) | `/api/v1/datasets` 未认证返回 `Authorization can't be empty`；无效 Key 返回 `code:109 API key is invalid!`；前端 umi.js 大量 `ragflow.io/docs` 链接 |")
md.append("| cas-func | **CAS 统一认证** | http://cas-func.ibosssoft.com.cn/cas | TGT/TGC 会话 | 前端 env.js `CAS_BASE_URL` 显式配置 |")
md.append("")
md.append("> **重要更正**: 之前从 ai-func JS 提取的 16 个 `/rag/pipelines/*` 接口在 ai-func 主域**全部404**——")
md.append("> 它们属于 rag-func(RagFlow) 服务的概念，但真实路径是 `/api/v1/*` 形式，并非 `/rag/pipelines/*`。")

md.append("")
md.append("## 二、验证统计")
md.append("")
from collections import Counter
vc = Counter(r[4] for r in svc_rows)
md.append(f"**接口总数**: {len(svc_rows)}")
md.append("")
md.append("| 验证状态 | 数量 | 说明 |")
md.append("|---|---|---|")
for k, n in vc.most_common():
    md.append(f"| {k} | {n} | |")
md.append("")

# 各服务统计
md.append("## 三、各服务接口明细")
md.append("")
for svc in ["dify_console", "dify_runtime", "ragflow", "cas"]:
    rows = [r for r in svc_rows if r[0] == svc]
    valid = sum(1 for r in rows if r[4].startswith("VALID"))
    md.append(f"### {svc}（{len(rows)} 个，有效 {valid}）")
    md.append("")
    md.append("| 路径 | 方法 | 验证 | 说明 |")
    md.append("|---|---|---|---|")
    for _, g, p, m, v, note in sorted(rows, key=lambda x: (x[2], x[3])):
        badge = {"VALID": "✅", "METHOD": "⚠️", "DISABLED": "🚫",
                 "OPEN": "🟢", "VALID-实测401": "✅", "VALID-代码探测": "🔧"}.get(v, "❓")
        md.append(f"| `{p}` | {m} | {badge} {v} | {note} |")
    md.append("")

md.append("## 四、验证方法说明")
md.append("")
md.append("1. **Dify 接口**: 对全部候选路径发送 GET/POST/PUT/PATCH/DELETE 探测请求，")
md.append("   根据 `401=存在需认证` / `405=存在方法错` / `503=被禁用` / `404=不存在` 判定。")
md.append("2. **RagFlow 接口**: 实测 `/api/v1/datasets`、`/api/v1/chats`、`/api/v1/agents` 存在，")
md.append("   其余按 RagFlow 官方 OpenAPI 结构补充（标注\"代码探测\"，需登录确认）。")
md.append("3. **局限**: 未登录态无法验证业务级参数/权限；接口的实际请求/响应 Schema 需有效 Token 后抓包补全。")
md.append("")
md.append("## 五、发现的接口问题")
md.append("")
md.append("| 问题 | 详情 |")
md.append("|---|---|")
md.append("| 伪路径（已剔除） | `/rag/pipelines/*`(16)，`/console/api/messages/*`(2)，`/saved-messages`(2)，`/datasets/create`，`/datasets/{id}/hitTesting`(大小写错误) |")
md.append("| 方法待修正 | `/workspaces/current`(PATCH→GET), `/members`(POST→GET), `/plugin/debugging-key`, `/plugin/install/pkg`, `/plugin/list/installations/ids`, `/tool-provider/mcp/update/{id}` |")
md.append("| 被禁用 | `/workspaces`、`/members/*/owner-transfer*`(3)、`/setup` 等 503「功能已禁用」|")
md.append("| 双 API 服务 | main 域 `/console/api/*` 与 `/runtime/console/api/*` 是不同后端（401 消息不同）|")

OUT_MD.write_text("\n".join(md), encoding="utf-8")

# ---------- Excel 报告 ----------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = "接口验证"
hdrs = ["服务", "分组", "接口路径", "HTTP方法", "验证状态", "说明"]
for c, h in enumerate(hdrs, 1):
    cell = ws.cell(1, c, h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center")

fills = {"VALID": "C6EFCE", "METHOD": "FFEB9C", "DISABLED": "F8CBAD",
         "OPEN": "C6EFCE", "VALID-实测401": "C6EFCE", "VALID-代码探测": "DDEBF7"}
for r, (svc, g, p, m, v, note) in enumerate(sorted(svc_rows, key=lambda x: (x[0], x[2])), 2):
    for c, val in enumerate([svc, g, p, m, v, note], 1):
        cell = ws.cell(r, c, val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    fill = fills.get(v)
    if fill:
        for c in range(1, 7):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{len(svc_rows)+1}"
for i, w in enumerate([16, 14, 60, 20, 18, 40], 1):
    from openpyxl.utils import get_column_letter
    ws.column_dimensions[get_column_letter(i)].width = w

# 汇总 sheet
ws2 = wb.create_sheet("汇总")
ws2["A1"] = "接口验证汇总"
ws2["A1"].font = Font(bold=True, size=14)
r = 3
for k, n in vc.most_common():
    ws2.cell(r, 1, k); ws2.cell(r, 2, n)
    r += 1
ws2.cell(r + 1, 1, "总计"); ws2.cell(r + 1, 2, len(svc_rows)).font = Font(bold=True)

wb.save(OUT_XLSX)

log(f"Markdown: {OUT_MD}")
log(f"Excel:    {OUT_XLSX}")
log(f"接口总数: {len(svc_rows)}")
for k, n in vc.most_common():
    log(f"  {k}: {n}")
print("\n".join(lines))