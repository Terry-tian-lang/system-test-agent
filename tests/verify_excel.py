#!/usr/bin/env python3
"""验证生成的 Excel 内容质量"""
import sys, glob, os
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

OUT = ROOT / "output" / "verify_excel.txt"
lines = []
def log(m=""):
    lines.append(m)

from openpyxl import load_workbook

files = sorted(glob.glob(str(ROOT / "output" / "*测试用例*.xlsx")), key=os.path.getmtime)
if not files:
    log("未找到测试用例 Excel")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    sys.exit(1)

f = files[-1]
log(f"FILE: {f}")
log(f"SIZE: {os.path.getsize(f)} bytes")

wb = load_workbook(f)
log(f"SHEETS: {wb.sheetnames}")
ws = wb["测试用例"]
log(f"DIMS: {ws.max_row - 1} cases x {ws.max_column} cols")
log("")

hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
log(f"HEADERS: {hdr}")

required = ["序号", "用例编号", "模块", "优先级", "类型", "用例标题",
            "前置条件", "操作步骤", "预期结果", "实际结果", "备注", "关联需求"]
log(f"匹配要求的12列: {hdr == required}")
log("")

# 完整性检查
empty_stats = {h: 0 for h in required}
for r in range(2, ws.max_row + 1):
    for c, h in enumerate(required, 1):
        v = ws.cell(r, c).value
        if v is None or str(v).strip() == "":
            empty_stats[h] += 1

log("=== 空值统计 (实际结果应为空，其他应填满) ===")
for h, n in empty_stats.items():
    flag = "OK(执行时填写)" if h == "实际结果" else ("!! 有空值" if n else "OK")
    log(f"  {h}: {n} 空  {flag}")
log("")

# 抽样展示
log("=== 抽样用例 ===")
for r in [2, 24, 47, 70, ws.max_row]:
    if r > ws.max_row or r < 2:
        continue
    log(f"--- 第 {r} 行 ---")
    for c, h in enumerate(required, 1):
        v = str(ws.cell(r, c).value or "").replace("\n", " ⏎ ")
        log(f"  {h}: {v[:180]}")
    log("")

# 摘要 sheet
if "测试摘要" in wb.sheetnames:
    ws2 = wb["测试摘要"]
    log("=== 测试摘要 Sheet ===")
    for r in range(1, min(ws2.max_row + 1, 45)):
        vals = [str(ws2.cell(r, c).value or "") for c in range(1, 4)]
        if any(v for v in vals):
            log("  " + " | ".join(v for v in vals if v))

# 样式检查
log("")
log("=== 样式检查 ===")
log(f"  冻结窗格: {ws.freeze_panes}")
log(f"  自动筛选: {ws.auto_filter.ref}")
log(f"  数据验证规则数: {len(ws.data_validations.dataValidation)}")
p0cell = None
for r in range(2, ws.max_row + 1):
    if str(ws.cell(r, 4).value).upper() == "P0":
        p0cell = ws.cell(r, 4)
        break
if p0cell:
    log(f"  P0单元格填充色: {p0cell.fill.start_color.rgb}")
    log(f"  P0单元格加粗: {p0cell.font.bold}")

OUT.write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines))
