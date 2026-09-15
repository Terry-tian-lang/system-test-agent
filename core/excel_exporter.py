"""
Excel 导出模块 - 将测试用例导出为格式化的 Excel 文件
"""
import os
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from core.logger import logger

from config.loader import get_config
from core.testcase_generator import TestCase


class ExcelExporter:
    """测试用例 Excel 导出器"""

    # 优先级颜色映射
    PRIORITY_COLORS = {
        "P0": "FF4444",   # 红色
        "P1": "FF9900",   # 橙色
        "P2": "FFDD00",   # 黄色
        "P3": "99CC00",   # 绿色
    }

    def __init__(self, config=None):
        self.config = config or get_config()
        self.excel_config = self.config.get_excel_config()

    def export(self, test_cases: List[TestCase],
               output_path: str = None,
               project_name: str = "系统测试") -> str:
        """
        将测试用例导出到 Excel 文件
        :param test_cases: 测试用例列表
        :param output_path: 输出文件路径（可选，自动生成）
        :param project_name: 项目名称
        :return: 输出文件路径
        """
        if not test_cases:
            logger.warning("没有测试用例数据，跳过导出")
            return ""

        # 确定输出路径
        if output_path is None:
            output_dir = self.excel_config.get("output_dir", "./output")
            os.makedirs(output_dir, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{project_name}_测试用例_{timestamp}.xlsx"
            output_path = os.path.join(output_dir, filename)

        logger.info(f"开始导出 {len(test_cases)} 条测试用例到 Excel...")

        wb = Workbook()
        ws = wb.active
        ws.title = "测试用例"

        # 1. 写入表头
        self._write_headers(ws)

        # 2. 写入数据
        self._write_data(ws, test_cases)

        # 3. 设置列宽
        self._set_column_widths(ws)

        # 4. 设置样式
        self._apply_styles(ws, len(test_cases))

        # 5. 添加数据验证
        self._add_validations(ws, len(test_cases))

        # 6. 添加摘要 Sheet
        self._add_summary_sheet(wb, test_cases, project_name)

        # 7. 保存文件
        os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
        wb.save(output_path)
        logger.info(f"Excel 文件已保存: {output_path}")

        return output_path

    def _write_headers(self, ws):
        """写入表头行"""
        headers = self.excel_config.get("headers", [
            "序号", "用例编号", "模块", "优先级", "类型",
            "用例标题", "前置条件", "操作步骤", "预期结果",
            "实际结果", "备注", "关联需求"
        ])
        style = self.excel_config.get("style", {})

        header_font = Font(
            bold=style.get("header_font_bold", True),
            color=style.get("header_font_color", "FFFFFF"),
            size=11,
        )
        header_fill = PatternFill(
            start_color=style.get("header_fill", "4472C4"),
            end_color=style.get("header_fill", "4472C4"),
            fill_type="solid"
        )
        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

    def _write_data(self, ws, test_cases: List[TestCase]):
        """写入测试用例数据"""
        for row_idx, case in enumerate(test_cases, 2):
            row_data = case.to_excel_row()
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    def _set_column_widths(self, ws):
        """设置列宽"""
        widths = self.excel_config.get("column_widths", {})
        headers = self.excel_config.get("headers", [
            "序号", "用例编号", "模块", "优先级", "类型",
            "用例标题", "前置条件", "操作步骤", "预期结果",
            "实际结果", "备注", "关联需求"
        ])

        default_widths = {
            "序号": 8, "用例编号": 15, "模块": 15, "优先级": 8,
            "类型": 12, "用例标题": 40, "前置条件": 30,
            "操作步骤": 50, "预期结果": 40, "实际结果": 20,
            "备注": 25, "关联需求": 15
        }

        for col_idx, header in enumerate(headers, 1):
            width = widths.get(header, default_widths.get(header, 15))
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _apply_styles(self, ws, row_count: int):
        """应用单元格样式"""
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for row in range(1, row_count + 2):
            for col in range(1, 13):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border

                # 优先级列着色 (第4列)
                if col == 4 and row > 1:
                    priority = str(cell.value or "").upper()
                    if priority in self.PRIORITY_COLORS:
                        cell.fill = PatternFill(
                            start_color=self.PRIORITY_COLORS[priority],
                            end_color=self.PRIORITY_COLORS[priority],
                            fill_type="solid"
                        )
                        cell.font = Font(bold=True, color="FFFFFF" if priority == "P0" else "000000")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # 序号列居中
                if col == 1 and row > 1:
                    cell.alignment = Alignment(horizontal="center", vertical="top")

        # 冻结首行
        style = self.excel_config.get("style", {})
        if style.get("freeze_panes"):
            ws.freeze_panes = style["freeze_panes"]

        # 自动筛选
        if style.get("auto_filter", True):
            ws.auto_filter.ref = f"A1:L{row_count + 1}"

    def _add_validations(self, ws, row_count: int):
        """添加数据验证（下拉列表）"""
        # 优先级下拉
        priority_dv = DataValidation(
            type="list",
            formula1='"P0,P1,P2,P3"',
            allow_blank=True
        )
        priority_dv.error = "请选择有效的优先级"
        priority_dv.errorTitle = "无效优先级"
        ws.add_data_validation(priority_dv)
        priority_dv.add(f"D2:D{row_count + 1}")

        # 类型下拉
        types = self.config.get_testcase_config().get("case_types", [
            "功能测试", "边界值测试", "异常测试", "兼容性测试", "安全测试", "性能相关"
        ])
        type_formula = '"' + ",".join(types) + '"'
        type_dv = DataValidation(
            type="list",
            formula1=type_formula,
            allow_blank=True
        )
        type_dv.error = "请选择有效的测试类型"
        type_dv.errorTitle = "无效类型"
        ws.add_data_validation(type_dv)
        type_dv.add(f"E2:E{row_count + 1}")

    def _add_summary_sheet(self, wb, test_cases: List[TestCase],
                           project_name: str):
        """添加摘要统计 Sheet"""
        ws = wb.create_sheet(title="测试摘要")

        # 标题
        ws["A1"] = f"{project_name} - 测试用例摘要"
        ws["A1"].font = Font(bold=True, size=14, color="4472C4")
        ws.merge_cells("A1:D1")

        # 基本信息
        ws["A3"] = "项目信息"
        ws["A3"].font = Font(bold=True, size=12)
        ws["A4"] = "项目名称"
        ws["B4"] = project_name
        ws["A5"] = "生成时间"
        ws["B5"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws["A6"] = "总用例数"
        ws["B6"] = len(test_cases)

        # 按优先级统计
        ws["A8"] = "优先级统计"
        ws["A8"].font = Font(bold=True, size=12)
        ws["A9"] = "优先级"
        ws["B9"] = "数量"
        ws["C9"] = "占比"
        for cell in [ws["A9"], ws["B9"], ws["C9"]]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

        priority_counts = {}
        for case in test_cases:
            p = case.优先级
            priority_counts[p] = priority_counts.get(p, 0) + 1

        row = 10
        for p in ["P0", "P1", "P2", "P3"]:
            count = priority_counts.get(p, 0)
            ws.cell(row=row, column=1, value=p)
            ws.cell(row=row, column=2, value=count)
            ws.cell(row=row, column=3, value=f"{count/len(test_cases)*100:.1f}%")
            row += 1

        # 按模块统计
        row += 1
        ws.cell(row=row, column=1, value="模块统计").font = Font(bold=True, size=12)
        row += 1
        ws.cell(row=row, column=1, value="模块名称").font = Font(bold=True)
        ws.cell(row=row, column=2, value="用例数").font = Font(bold=True)
        ws.cell(row=row, column=3, value="占比").font = Font(bold=True)
        for cell in [ws.cell(row=row, column=c) for c in [1, 2, 3]]:
            cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        row += 1

        module_counts = {}
        for case in test_cases:
            m = case.模块
            module_counts[m] = module_counts.get(m, 0) + 1

        for module, count in sorted(module_counts.items(), key=lambda x: -x[1]):
            ws.cell(row=row, column=1, value=module)
            ws.cell(row=row, column=2, value=count)
            ws.cell(row=row, column=3, value=f"{count/len(test_cases)*100:.1f}%")
            row += 1

        # 按类型统计
        row += 1
        ws.cell(row=row, column=1, value="测试类型统计").font = Font(bold=True, size=12)
        row += 1
        ws.cell(row=row, column=1, value="类型").font = Font(bold=True)
        ws.cell(row=row, column=2, value="数量").font = Font(bold=True)
        ws.cell(row=row, column=3, value="占比").font = Font(bold=True)
        for cell in [ws.cell(row=row, column=c) for c in [1, 2, 3]]:
            cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        row += 1

        type_counts = {}
        for case in test_cases:
            t = case.类型
            type_counts[t] = type_counts.get(t, 0) + 1

        for t, count in sorted(type_counts.items(), key=lambda x: -x[1]):
            ws.cell(row=row, column=1, value=t)
            ws.cell(row=row, column=2, value=count)
            ws.cell(row=row, column=3, value=f"{count/len(test_cases)*100:.1f}%")
            row += 1

        # 设置列宽
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 15


def export_to_excel(test_cases: List[TestCase],
                    output_path: str = None,
                    project_name: str = "系统测试") -> str:
    """
    便捷函数：导出测试用例到 Excel
    """
    exporter = ExcelExporter()
    return exporter.export(test_cases, output_path, project_name)
